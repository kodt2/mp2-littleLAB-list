import openpyxl
import datetime
import fpdf

def nedeed_text1_names(pdf,KNV, K, Z_name, VAE, B_name, predsed_name, predsed_phone):
    pdf.cell(200, 13, txt="", ln=1, align="R")
    pdf.cell(190, 6, txt="Начальнику управления безопасности", ln=1, align="R")
    pdf.cell(190, 6, txt="ННГУ им. Н.И. Лобачевского", ln=1, align="R")
    pdf.cell(190, 6, txt=VAE, ln=1, align="R")
    pdf.cell(190, 6, txt="от заместителя директора по учебно-воспитательной работе", ln=1, align="R")
    pdf.cell(190, 6, txt="Института Информационных", ln=1, align="R")
    pdf.cell(190, 6, txt="Технологий, Математики и Механики", ln=1, align="R")
    pdf.cell(190, 6, txt=KNV, ln=1, align="R")
    pdf.cell(190, 8, txt="", ln=1, align="C")
    pdf.cell(190, 15, txt="Служебная записка.", ln=1, align="C")
    pdf.cell(10)
    pdf.multi_cell(180, 6, txt="            Институт Информационных технологий математики"
                                   " и механики просит\n "
                                   "Вас предоставить доступ в "
                                   "аудитории корпуса № " + '2' + " ННГУ им. Н.И. Лобачевского "
                                                                   "для проведения мероприятия Студенческого "
                                                                   "Совета ИИТММ, следующим обучающимся:",
                   align="L")
    pdf.cell(190, 6, txt="", ln=1, align="l")
    return pdf

def nedeed_text2(pdf,KaNV, KNV, Z_name, VAE, B_name, predsed_name, predsed_phone):
    pdf.cell(10)
    pdf.cell(190, 6, txt="Ответственные:", ln=1, align="l")
    pdf.cell(10)
    pdf.cell(190, 6, txt=Z_name+" "+B_name+" "+ KNV, ln=1, align="l")
    pdf.cell(190, 6, txt="", ln=1, align="l")
    pdf.cell(10)
    pdf.cell(100, 6, txt="Заместитель директора ИИТММ", ln=1, align="l")
    pdf.cell(10)
    pdf.cell(100, 6, txt="по учебно-воспитательной работе", align="l")
    pdf.cell(40)
    pdf.cell(90, 6, txt=KNV, ln=1, align="r")
    pdf.cell(190, 6, txt="", ln=1, align="l")
    pdf.cell(190, 6, txt="Председатель СС ИИТММ:"+predsed_name+" "+predsed_phone, ln=1, align="R")
    return pdf


def list_of_names(pdf, arr, start, end):
    def add_name_cell(pdf, index, name):
        pdf.cell(10)
        pdf.cell(190, 6, txt=f"{index + 1}. {name}", ln=1, align="l")

    for i in range(start, end):
        add_name_cell(pdf, i, arr[i])

    pdf.cell(190, 6 * (16 - end + start), txt="", ln=1, align="C")
    return pdf


def setup_pdf():
    npdf = fpdf.FPDF(format='letter')
    npdf.add_page()
    npdf.add_font('Times', '', 'times.ttf', uni=True)
    npdf.set_font("Times", size=14)
    return npdf

def work_names(arr_of_names, KaNV, KNV,  Z_name, VAE, B_name, predsed_name, predsed_phone):
    npdf = setup_pdf()
    count_pages = (len(arr_of_names) + 13) // 14  # количество страниц

    for page_number in range(count_pages):
        start = page_number * 14
        end = min(start + 14, len(arr_of_names))

        npdf = nedeed_text1_names(npdf, KaNV, KNV, Z_name, VAE, B_name, predsed_name, predsed_phone)
        npdf = list_of_names(npdf, arr_of_names, start, end)
        npdf = nedeed_text2(npdf, KaNV, KNV, Z_name, VAE, B_name, predsed_name, predsed_phone)

        if page_number < count_pages - 1:
            npdf.add_page()

    npdf.output("names.pdf")

